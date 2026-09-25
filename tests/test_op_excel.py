"""OP en Excel real (COM): formato uniforme, editable, casilla de ciudad y vendedor."""

import datetime as dt

import openpyxl
import pytest
import win32com.client as win32

from virlan_bot import config, op_filler
from virlan_bot.models import LineaRenovacion

pytestmark = [pytest.mark.excel, pytest.mark.machote]

CIUDAD_LARGA = "Ciudad Nezahualcoyotl de los Reyes de la Paz Grande"


def linea(tel, modelo, plan="ATT Armalo Negocios 599 CPP CTRL", extra="CTRL", costo=None):
    return LineaRenovacion(
        telefono=tel, plan_tarifario=plan, plazo_meses="24", modelo=modelo,
        marca_modelo_color=f"{modelo}, NEGRO", addon_extra=extra, costo_equipo=costo,
    )


@pytest.fixture(scope="module")
def op(tmp_path_factory):
    """Una sola corrida de Excel para todas las aserciones del módulo."""
    from tests.conftest import _fila_precios  # noqa: F401  (los catálogos se arman aquí abajo)
    import csv

    base = tmp_path_factory.mktemp("op")
    # catálogos sintéticos mínimos
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Precios"
    ws.cell(4, 5, "APPLE IPHONE 17 256GB 5G"); ws.cell(4, 6, 19999); ws.cell(4, 24, 13199)
    ws.cell(5, 5, "SAMSUNG GALAXY ZFOLD 8 ULTRA 512GB 5G"); ws.cell(5, 6, 53499); ws.cell(5, 24, 37759)
    for r in (4, 5):
        ws.cell(r, 42, dt.date(2026, 1, 1)); ws.cell(r, 43, dt.date(2050, 1, 1))
    lista = base / "lista.xlsx"; wb.save(lista)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Master"
    ws.append(["Marca", "Modelo", "Precio inicial", "Precio Base"])
    ws.append(["APPLE", "APPLE IPHONE 17 256GB", 19999, 19999])
    ws.append(["SAMSUNG", "SAMSUNG GALAXY Z FOLD8 ULTRA", 48499, 48499])
    mpe = base / "mpe.xlsx"; wb.save(mpe)
    ladas = base / "ladas.csv"
    with open(ladas, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f); w.writerow(["Lada", "Estado", "Ciudad_Region_Principal", "Tipo", "C", "N"])
        w.writerow(["55", "Ciudad de Mexico", "Ciudad de Mexico", "M", "Si", ""])
        w.writerow(["221", "Puebla", "San Martin Texmelucan (zona)", "R", "Si", ""])
        w.writerow(["559", "Estado de Mexico", CIUDAD_LARGA, "R", "Si", ""])   # nombre larguísimo a propósito

    from tests.conftest import cliente as _cliente_fixture  # noqa: F401
    from virlan_bot.models import ClienteContrato

    c = ClienteContrato(
        razon_social="EMPRESA PRUEBA SA DE CV", numero_cuenta="500000001", representante_legal="JUAN PEREZ LOPEZ",
        domicilio_calle="CALLE UNO", domicilio_numero="12", domicilio_colonia="COL CENTRO",
        domicilio_ciudad="CUAUHTEMOC", domicilio_estado="CIUDAD DE MEXICO", domicilio_cp="06000",
        telefono="5512345678", rfc="EPR010101AB1", tipo_identificacion="INE", numero_identificacion="123",
    )
    c.lineas = [
        linea("5511111111", "APPLE IPHONE 17 256GB 5G"),                          # CDMX, con precio
        linea("2215699113", "SAMSUNG GALAXY ZFOLD 8 ULTRA 512GB 5G"),             # Puebla; lista != Master
        linea("5599999999", "SAMSUNG GALAXY A14 128GB 4G"),                       # ciudad larguísima; equipo desconocido
        linea("NA", "SIM CARD", plan="ATT Ármalo Negocios $299 C"),               # SIM: N/A donde no aplica
    ]
    salida = base / "op.xlsx"
    # La regla vigente reduce las ciudades a su primera palabra; para ejercitar la reducción de
    # letra por casilla se fuerza un nombre completo larguísimo en UNA línea.
    original = op_filler.ciudad_dn_por_telefono
    mp = pytest.MonkeyPatch()
    mp.setattr(op_filler, "ciudad_dn_por_telefono",
               lambda t, csv: CIUDAD_LARGA if t.endswith("9999") else original(t, csv))
    try:
        res = op_filler.llenar_op(
            config.MACHOTE_OP_XLSX, salida, c, ladas, "RENOVACION", lista, mpe, None,
            fecha_contratacion=dt.date(2026, 9, 24), ejecutivo="ALEJANDRO CASTERA ARELLANOS",
        )
    finally:
        mp.undo()
    return salida, res, openpyxl.load_workbook(salida).active


def test_ciudad_dn(op):
    _, _, ws = op
    assert [ws[f"D{r}"].value for r in (16, 17, 18, 19)] == ["CDMX", "Puebla", CIUDAD_LARGA, "N/A"]


def test_solo_la_casilla_de_ciudad_desbordada_reduce_su_letra(op):
    _, _, ws = op
    normales = {ws[f"D{r}"].font.sz for r in (16, 17, 19)}
    assert len(normales) == 1                          # las demás conservan el tamaño normal
    assert ws["D18"].font.sz < normales.pop()          # solo la desbordada baja su letra


def test_precios_mpe_y_sim(op):
    _, res, ws = op
    assert (ws["J16"].value, ws["K16"].value, ws["N16"].value) == (19999, 13199, 283.34)
    assert (ws["J17"].value, ws["K17"].value, ws["N17"].value) == (53499, 37759, 655.84)   # gana la lista
    assert (ws["J18"].value, ws["N18"].value) == ("N/A", "N/A")                              # sin sustituto
    assert [ws[f"{c}19"].value for c in "JKN"] == ["N/A", "N/A", "N/A"]                     # SIM
    assert (ws["P19"].value, ws["S19"].value, ws["Y19"].value) == ("X", 50, 349)
    texto = " ".join(res.alertas)
    assert "MISMO modelo" in texto and "Consultar por qué difieren" in texto


def test_todas_las_filas_con_el_mismo_formato_y_sin_reducir_hasta_ajustar(op):
    _, _, ws = op
    for col in "CEFHJKLNOPQSTUVXY":
        formatos = {(ws[f"{col}{r}"].font.sz, ws[f"{col}{r}"].alignment.horizontal) for r in (16, 17, 18, 19)}
        assert len(formatos) == 1, (col, formatos)
    assert not [c.coordinate for row in ws.iter_rows(min_row=1, max_row=40) for c in row
                if c.alignment and c.alignment.shrink_to_fit]


def test_fecha_y_encabezado(op):
    _, _, ws = op
    assert ws["X6"].value.date() == dt.date(2026, 9, 24)
    assert ws["C9"].value == "EMPRESA PRUEBA SA DE CV" and str(ws["J9"].value) == "500000001"
    assert ws["N12"].value == "INE //123"


def test_vendedor_y_firma_en_los_recuadros(op):
    """Los recuadros de texto (shapes) solo se pueden leer por COM."""
    salida, _, _ = op
    excel = win32.gencache.EnsureDispatch("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    try:
        wb = excel.Workbooks.Open(str(salida.resolve()))
        try:
            hoja = wb.Worksheets("OP_OK")
            assert hoja.Shapes("TextBox 13").TextFrame2.TextRange.Text == "ALEJANDRO CASTERA ARELLANOS"
            assert hoja.Shapes("TextBox 5").TextFrame2.TextRange.Text == "JUAN PEREZ LOPEZ"
            assert hoja.Shapes("Rectangle 3").TextFrame2.TextRange.Text == "X"        # RENOVACIÓN
            assert hoja.Shapes("Rectangle 1").TextFrame2.TextRange.Text == ""
        finally:
            wb.Close(False)
    finally:
        excel.Quit()
