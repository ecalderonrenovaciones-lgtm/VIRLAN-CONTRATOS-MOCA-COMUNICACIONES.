"""SAE, control de renovación y LAYOUT DE VINCULACION (archivos sintéticos)."""

import openpyxl
import pytest

from virlan_bot.models import ClienteContrato
from virlan_bot.renovacion_extractor import extraer_lineas_renovacion
from virlan_bot.vinculacion_extractor import extraer_vinculacion

CUENTA = "500000001"


def xlsx(ruta, fila_encabezado, encabezados, filas):
    wb = openpyxl.Workbook()
    ws = wb.active
    for c, h in enumerate(encabezados, 1):
        ws.cell(fila_encabezado, c, h)
    for i, fila in enumerate(filas, fila_encabezado + 1):
        for c, v in enumerate(fila, 1):
            ws.cell(i, c, v)
    wb.save(ruta)
    return ruta


ENC_CONTROL = ["No", "EMPRESA", "Cuenta/Sub cuenta", "Región", "Línea de Renovar", "Plan", "Modelo", "Color", "Plazo", "Costo de Equipo"]
ENC_SAE = ["No", "Empresa", "Cuenta / Sub Cuenta", "REDES SOCIALES", "STREAMING", "Addon Extra", "Líneas a Renovar", "Modelo",
           "Color", "Plan", "Plazo", "Modalidad      MPP / CPP"]   # espacios de largo variable, como el real


def control(tmp_path, lineas):
    """lineas: [(tel, plan, modelo, color, plazo, costo)]"""
    return xlsx(tmp_path / "control.xlsx", 7, ENC_CONTROL,
                [[i + 1, "E", CUENTA, "CENTRO", t, p, m, c, pl, co] for i, (t, p, m, c, pl, co) in enumerate(lineas)])


def sae(tmp_path, lineas, extra="CTRL", modalidad="CPP"):
    return xlsx(tmp_path / "sae.xlsx", 8, ENC_SAE,
                [[i + 1, "E", CUENTA, "Facebook", "Spotify", extra, t, m, c, p, f"{pl} MESES", modalidad]
                 for i, (t, p, m, c, pl, co) in enumerate(lineas)])


L1 = ("5511111111", "ATT Armalo Negocios 599 CPP CTRL", "APPLE IPHONE 17 256GB 5G", "NEGRO", "24", 13199)
L2 = ("5522222222", "ATT Armalo Negocios 299 CPP LIBRE", "HONOR X5D 128GB 4G", "NEGRO", "24", 1259)


def cliente_vacio():
    return ClienteContrato(numero_cuenta=CUENTA)


def test_sae_es_la_fuente_principal_y_trae_addon_y_modalidad(tmp_path):
    c = cliente_vacio()
    extraer_lineas_renovacion(c, control(tmp_path, [L1, L2]), sae(tmp_path, [L1, L2]))
    assert [l.telefono for l in c.lineas] == ["5511111111", "5522222222"]
    assert c.lineas[0].plazo_meses == "24"                    # 'MESES' se quita
    assert c.lineas[0].addon_extra == "CTRL" and c.lineas[0].modalidad_mpp_cpp == "CPP"
    assert c.alertas == []


def test_costo_de_equipo_sigue_al_telefono_aunque_el_orden_difiera(tmp_path):
    """Regresión: el traspaso por posición asignaba el costo a la línea equivocada."""
    c = cliente_vacio()
    extraer_lineas_renovacion(c, control(tmp_path, [L2, L1]), sae(tmp_path, [L1, L2]))
    costos = {l.telefono: l.costo_equipo for l in c.lineas}
    assert costos == {"5511111111": 13199.0, "5522222222": 1259.0}


def test_lineas_sim_sin_telefono_se_emparejan_por_posicion(tmp_path):
    sim = ("NA", "ATT Ármalo Negocios $299 C", "SIM CARD", "NA", "24", 0)
    c = cliente_vacio()
    extraer_lineas_renovacion(c, control(tmp_path, [L1, sim, sim]), sae(tmp_path, [L1, sim, sim]))
    assert len(c.lineas) == 3 and c.lineas[1].costo_equipo == 0.0
    assert c.lineas[1].marca_modelo_color == "SIM CARD"      # el color 'NA' no se concatena


def test_filas_sin_telefono_al_final_de_la_hoja_no_son_lineas(tmp_path):
    ruta = control(tmp_path, [L1])
    wb = openpyxl.load_workbook(ruta)
    ws = wb.active
    ws.cell(1048576, 2, "E"); ws.cell(1048576, 3, CUENTA)      # fila suelta con la cuenta pero sin teléfono
    wb.save(ruta)
    c = cliente_vacio()
    extraer_lineas_renovacion(c, ruta, sae(tmp_path, [L1]))
    assert len(c.lineas) == 1 and not any("línea(s)" in a for a in c.alertas)


def test_solo_control_o_solo_sae(tmp_path):
    a, b = cliente_vacio(), cliente_vacio()
    extraer_lineas_renovacion(a, control(tmp_path, [L1]), None)
    extraer_lineas_renovacion(b, None, sae(tmp_path, [L1]))
    assert a.lineas[0].costo_equipo == 13199.0 and b.lineas[0].addon_extra == "CTRL"


def test_conteos_o_planes_distintos_entre_archivos_se_avisan(tmp_path):
    c = cliente_vacio()
    otro_plan = ("5511111111", "ATT Armalo Negocios 799 CPP CTRL", "APPLE IPHONE 17 256GB 5G", "NEGRO", "24", 13199)
    extraer_lineas_renovacion(c, control(tmp_path, [L1, L2]), sae(tmp_path, [otro_plan]))
    assert any("línea(s)" in a for a in c.alertas) and any("Discrepancia" in a for a in c.alertas)


def test_cuenta_sin_lineas_avisa(tmp_path):
    c = ClienteContrato(numero_cuenta="999999999")
    extraer_lineas_renovacion(c, control(tmp_path, [L1]), sae(tmp_path, [L1]))
    assert c.lineas == [] and "lineas" in c.campos_faltantes and c.alertas


# ------------------------------------------------------------------ vinculación
ENC_LAYOUT = ["Consecutivo", "ATTUID", "Tipo Movimiento", "# cuenta AMDOCS", "RFC", "CURP", "Tipo identificación", "# de identificación"]


def layout(tmp_path, rfc="EPR010101AB1", cuenta=CUENTA, tipo="INE", numero="1234567890"):
    return xlsx(tmp_path / "layout.xlsx", 1, ENC_LAYOUT, [[0, "X", "RENOVACION", cuenta, rfc, "CURP", tipo, numero]])


def test_vinculacion_llena_rfc_e_identificacion(tmp_path):
    c = cliente_vacio()
    extraer_vinculacion(c, layout(tmp_path))
    assert (c.rfc, c.tipo_identificacion, c.numero_identificacion) == ("EPR010101AB1", "INE", "1234567890")
    assert c.campos_faltantes == []


def test_vinculacion_pisa_el_rfc_de_la_ficha_con_aviso(tmp_path):
    c = cliente_vacio()
    c.rfc = "OTRO010101XX0"
    extraer_vinculacion(c, layout(tmp_path))
    assert c.rfc == "EPR010101AB1" and any("no coincide" in a for a in c.alertas)


def test_vinculacion_sin_la_cuenta_marca_faltantes(tmp_path):
    c = cliente_vacio()
    extraer_vinculacion(c, layout(tmp_path, cuenta="111111111"))
    assert {"rfc", "tipo_identificacion", "numero_identificacion"} <= set(c.campos_faltantes) and c.alertas


def test_vinculacion_sin_rfc_no_inventa(tmp_path):
    c = cliente_vacio()
    extraer_vinculacion(c, layout(tmp_path, rfc=None))
    assert c.rfc == "" and "rfc" in c.campos_faltantes
