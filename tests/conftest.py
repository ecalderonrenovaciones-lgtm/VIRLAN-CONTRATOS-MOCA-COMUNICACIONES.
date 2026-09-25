"""Fixtures compartidas. Todo es SINTÉTICO (nombres, cuentas y RFC inventados):
las pruebas no dependen de ni exponen datos reales de clientes."""

from __future__ import annotations

import csv
import datetime as dt
import zipfile
from pathlib import Path

import openpyxl
import pytest

from virlan_bot import config
from virlan_bot.models import ClienteContrato, LineaRenovacion

MACHOTE_CONTRATO = config.MACHOTE_CONTRATO_PDF
MACHOTE_OP = config.MACHOTE_OP_XLSX


def pytest_collection_modifyitems(config, items):  # noqa: ARG001
    """Omite las pruebas marcadas si el recurso externo no está disponible."""
    import importlib.util

    hay_machotes = MACHOTE_CONTRATO.exists() and MACHOTE_OP.exists()
    hay_excel = importlib.util.find_spec("win32com") is not None
    for item in items:
        if "machote" in item.keywords and not hay_machotes:
            item.add_marker(pytest.mark.skip(reason="faltan los machotes en ARCHIVOS INTERCAMBIABLES"))
        if "excel" in item.keywords and not (hay_machotes and hay_excel):
            item.add_marker(pytest.mark.skip(reason="requiere Excel/pywin32 y los machotes"))


# ------------------------------------------------------------------ fichas .docx
def crear_docx(ruta: Path, parrafos: list[str]) -> Path:
    """Word mínimo (solo word/document.xml), suficiente para ficha_extractor."""
    cuerpo = "".join(
        f"<w:p><w:r><w:t>{p}</w:t></w:r></w:p>" if p else "<w:p></w:p>" for p in parrafos
    )
    xml = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"<w:body>{cuerpo}</w:body></w:document>"
    )
    with zipfile.ZipFile(ruta, "w") as z:
        z.writestr("word/document.xml", xml)
    return ruta


DOM_FISCAL = "CALLE UNO, 12, COL CENTRO, CUAUHTEMOC, CIUDAD DE MEXICO, CIUDAD DE MEXICO, 06000, MEX"
DOM_ENTREGA = "AV DOS, 45 INT B, COL NORTE, GUADALAJARA, GUADALAJARA, JALISCO, 44100, MEX"


@pytest.fixture
def ficha_simple(tmp_path) -> Path:
    return crear_docx(
        tmp_path / "ficha.docx",
        ["EMPRESA PRUEBA SA DE CV", "500000001", "Sr. JUAN PEREZ LOPEZ", DOM_FISCAL,
         "EPR010101AB1", "5512345678", "juan@prueba.mx"],
    )


@pytest.fixture
def ficha_entrega(tmp_path) -> Path:
    """Formato vigente desde 2026-09-23: banner de envío + rótulos de bloque."""
    return crear_docx(
        tmp_path / "ficha_entrega.docx",
        ["EMPRESA PRUEBA SA DE CV", "500000001", "EL PAQUETE SE ENVIA A ESTE DOMICILIO:ENTREGA",
         "Sra. MARIA GOMEZ RUIZ", "DOMICILIO FISCAL", "", DOM_FISCAL, "EPR010101AB1", "5512345678",
         "maria@prueba.mx", "", "", "ENTREGA", "", DOM_ENTREGA],
    )


# ------------------------------------------------------------------ catálogos sintéticos
@pytest.fixture
def ladas_csv(tmp_path) -> Path:
    ruta = tmp_path / "ladas.csv"
    filas = [
        ("55", "Ciudad de Mexico", "Ciudad de Mexico", "Metropolitana"),
        ("55", "Estado de Mexico", "Ciudad de Mexico", "Metropolitana"),
        ("33", "Jalisco", "Guadalajara / Zapopan", "Metropolitana"),
        ("81", "Nuevo Leon", "Monterrey", "Metropolitana"),
        ("221", "Puebla", "San Martin Texmelucan (zona)", "Regional"),
        ("223", "Puebla", "San Martin Texmelucan / Tlaxcala", "Regional"),
        ("223", "Tlaxcala", "San Martin Texmelucan / Tlaxcala", "Regional"),
        ("722", "Estado de Mexico", "Estado de Mexico (Toluca)", "Regional"),
        ("999", "Yucatan", "Merida", "Metropolitana"),
        ("998", "Quintana Roo", "Cancun", "Metropolitana"),
        ("442", "Queretaro", "Queretaro", "Metropolitana"),
        ("444", "San Luis Potosi", "San Luis Potosi", "Metropolitana"),
    ]
    with open(ruta, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["Lada", "Estado", "Ciudad_Region_Principal", "Tipo", "Confirmado_en_IFT_2026", "Notas"])
        for lada, estado, ciudad, tipo in filas:
            w.writerow([lada, estado, ciudad, tipo, "Si", ""])
    return ruta


HOY = dt.date(2026, 9, 24)


def _fila_precios(ws, fila, familia, full, dif24_599=None, inicio=dt.date(2026, 1, 1), fin=dt.date(2050, 1, 1)):
    ws.cell(fila, 5, familia)
    ws.cell(fila, 6, full)
    ws.cell(fila, 42, inicio)
    ws.cell(fila, 43, fin)
    if dif24_599 is not None:
        # bloque del plan 599 = índice 3, plazo 24 meses = índice 2 -> columna 7 + 5*3 + 2 = 24
        ws.cell(fila, 24, dif24_599)


@pytest.fixture
def lista_precios(tmp_path) -> Path:
    """Hoja 'Precios': encabezado en fila 3, datos desde la 4 (estructura real)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Precios"
    _fila_precios(ws, 4, "APPLE IPHONE 17 256GB 5G", 19999, 13199)
    _fila_precios(ws, 5, "APPLE IPHONE 17 AIR 256GB 5G", 25999, 19299)
    _fila_precios(ws, 6, "APPLE IPHONE 17 PRO 256GB 5G", 28499, 22999)
    _fila_precios(ws, 7, "APPLE IPHONE 17 PRO MAX 256GB 5G", 30999, 25299)
    _fila_precios(ws, 8, "SAMSUNG GALAXY ZFOLD 8 ULTRA 512GB 5G", 53499, 37759)
    _fila_precios(ws, 9, "SAMSUNG GALAXY ZFOLD 8 512GB 5G", 48499, 33000)
    _fila_precios(ws, 10, "HONOR MAGIC 8 LITE BUNDLE", 9999, 2869)
    _fila_precios(ws, 11, "HONOR X5D 128GB 4G", 3999, "-")  # '-': no se ofrece en ese plan/plazo
    # dos vigencias del mismo equipo: solo debe contar la que cubre HOY
    _fila_precios(ws, 12, "OPPO A80 256GB 5G", 6999, 1599, dt.date(2026, 1, 1), dt.date(2026, 9, 20))
    _fila_precios(ws, 13, "OPPO A80 256GB 5G", 6999, 1199, dt.date(2026, 9, 21), dt.date(2050, 1, 1))
    ruta = tmp_path / "Lista de Precios TEST.xlsx"
    wb.save(ruta)
    return ruta


@pytest.fixture
def calculo_mpe(tmp_path) -> Path:
    """Hoja 'Master': B modelo, D precio base (encabezado en fila 1)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Master"
    ws.append(["Marca", "Modelo", "Precio inicial", "Precio Base"])
    filas = [
        ("APPLE", "APPLE IPHONE 17 256GB", 19999),          # sin '5G'
        ("APPLE", "APPLE IPHONE 17 PRO 256GB", 28499),
        ("SAMSUNG", "SAMSUNG GALAXY Z FOLD8", 32999),
        ("SAMSUNG", "SAMSUNG GALAXY Z FOLD8 ULTRA", 48499),  # precio distinto al de la lista
        ("HONOR", "HONOR MAGIC8 LITE BUNDLE", 9999),
        ("HONOR", "Honor X5d", 3999),                         # sin GB ni 4G
        ("OPPO", "OPPO FIND X8 PRO 512GB 5G", 24999),         # NO debe usarse para nada
    ]
    for marca, modelo, precio in filas:
        ws.append([marca, modelo, precio, precio])
    ruta = tmp_path / "Calculo de MPE TEST.xlsx"
    wb.save(ruta)
    return ruta


# ------------------------------------------------------------------ cliente de prueba
@pytest.fixture
def cliente() -> ClienteContrato:
    c = ClienteContrato(
        razon_social="EMPRESA PRUEBA SA DE CV",
        numero_cuenta="500000001",
        representante_legal="JUAN PEREZ LOPEZ",
        domicilio_calle="CALLE UNO",
        domicilio_numero="12",
        domicilio_colonia="COL CENTRO",
        domicilio_ciudad="CUAUHTEMOC",
        domicilio_municipio="CIUDAD DE MEXICO",
        domicilio_estado="CIUDAD DE MEXICO",
        domicilio_cp="06000",
        telefono="5512345678",
        correo="juan@prueba.mx",
        rfc="EPR010101AB1",
        tipo_identificacion="INE",
        numero_identificacion="1234567890",
    )
    c.lineas = [
        LineaRenovacion(
            telefono="5511111111", plan_tarifario="ATT Armalo Negocios 599 CPP CTRL", plazo_meses="24",
            modelo="APPLE IPHONE 17 256GB 5G", marca_modelo_color="APPLE IPHONE 17 256GB 5G, NEGRO",
            addon_extra="CTRL",
        )
    ]
    return c
