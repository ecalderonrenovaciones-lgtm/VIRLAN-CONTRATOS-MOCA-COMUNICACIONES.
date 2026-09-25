"""Punta a punta: correo sintético -> CONTRATO + OP reales (usa Excel por COM).

Todo el material es inventado (cuenta 500000001, nombres ficticios); los catálogos de
precios/MPE/ladas son los sintéticos de conftest.py. Cada corrida de Excel tarda ~15-25 s.
"""

import datetime as dt
import re
import time
from email.message import EmailMessage

import fitz
import openpyxl
import pytest

from tests.conftest import DOM_FISCAL, crear_docx
from virlan_bot import cli, config

pytestmark = [pytest.mark.excel, pytest.mark.machote]

FECHA = dt.date(2026, 9, 24)
PLAN_EQ = "ATT Armalo Negocios 599 CPP CTRL"
PLAN_SIM = "ATT Ármalo Negocios $299 C"


def _xlsx(ruta, fila_encabezado, encabezados, filas):
    wb = openpyxl.Workbook()
    ws = wb.active
    for c, h in enumerate(encabezados, 1):
        ws.cell(fila_encabezado, c, h)
    for i, fila in enumerate(filas, fila_encabezado + 1):
        for c, v in enumerate(fila, 1):
            ws.cell(i, c, v)
    wb.save(ruta)
    return ruta.read_bytes()


def crear_correo(tmp_path, nombre_correo, asunto, lineas, cuerpo="Pendiente de contrato"):
    """lineas: [(telefono, plan, modelo, color, costo)]; devuelve la ruta del .eml."""
    d = tmp_path / "adj"
    d.mkdir(exist_ok=True)
    cuenta = "500000001"
    control = _xlsx(
        d / "c.xlsx", 7,
        ["No", "EMPRESA", "Cuenta/Sub cuenta", "Región", "Línea de Renovar", "Plan", "Modelo", "Color", "Plazo", "Costo de Equipo"],
        [[i + 1, "EMPRESA PRUEBA", cuenta, "CENTRO", t, p, m, c, "24", costo] for i, (t, p, m, c, costo) in enumerate(lineas)],
    )
    sae = _xlsx(
        d / "s.xlsx", 8,
        ["No", "Empresa", "Cuenta / Sub Cuenta", "REDES SOCIALES", "STREAMING", "Addon Extra", "Líneas a Renovar",
         "Modelo", "Color", "Plan", "Plazo", "Modalidad      MPP / CPP"],
        [[i + 1, "EMPRESA PRUEBA", cuenta, "Facebook; Gmail", "Spotify", "CTRL", t, m, c, p, "24 MESES", "CPP"]
         for i, (t, p, m, c, costo) in enumerate(lineas)],
    )
    layout = _xlsx(
        d / "l.xlsx", 1, ["Consecutivo", "ATTUID", "Tipo Movimiento", "# cuenta AMDOCS", "RFC", "CURP", "Tipo identificación",
                          "# de identificación", "Fecha de Validacion"],
        [[0, "X", "RENOVACION", cuenta, "EPR010101AB1", "XXXX010101HDFXXX00", "INE", "1234567890", 46288]],
    )
    ficha = crear_docx(d / "f.docx", ["EMPRESA PRUEBA SA DE CV", cuenta, "Sr. JUAN PEREZ LOPEZ", DOM_FISCAL,
                                      "EPR010101AB1", "5512345678", "juan@prueba.mx"]).read_bytes()
    doc = fitz.open()
    doc.new_page().insert_text((72, 72), "Cotejado contra original 24-09-2026")
    ine = doc.tobytes()

    msg = EmailMessage()
    msg["Subject"] = asunto
    msg.set_content(cuerpo)
    for nombre, datos, sub in (
        ("01control de renovacion.xlsx", control, "vnd.ms-excel"), ("03SAE 2.2.xlsx", sae, "vnd.ms-excel"),
        ("02Copia de LAYOUT DE VINCULACION EMPRESARIAL CURP.xlsx", layout, "vnd.ms-excel"),
        ("EMPRESA PRUEBA SA DE CV.docx", ficha, "msword"), ("INE_Sr. JUAN PEREZ LOPEZ.pdf", ine, "pdf"),
    ):
        msg.add_attachment(datos, maintype="application", subtype=sub, filename=nombre)
    ruta = tmp_path / nombre_correo
    ruta.write_bytes(bytes(msg))
    return ruta


@pytest.fixture
def entorno(tmp_path, monkeypatch, lista_precios, calculo_mpe, ladas_csv):
    monkeypatch.setattr(config, "SALIDA_DIR", tmp_path / "salida")
    monkeypatch.setattr(config, "CONTRATOS_TERMINADOS_DIR", tmp_path / "TERM")
    monkeypatch.setattr(config, "lista_precios_xlsx", lambda: lista_precios)
    monkeypatch.setattr(config, "calculo_mpe_xlsx", lambda: calculo_mpe)
    monkeypatch.setattr(config, "ladas_csv", lambda: ladas_csv)
    return tmp_path


EQUIPO = ("5511111111", PLAN_EQ, "APPLE IPHONE 17 256GB 5G", "NEGRO", 13199)
SIM1 = ("NA", PLAN_SIM, "SIM CARD", "NA", 0)
SIM2 = ("NA", PLAN_SIM, "SIM CARD", "NA", 0)


def texto(pdf):
    return re.sub(r"\s+", " ", fitz.open(pdf)[0].get_text())


def fila(ws, r):
    return {c: ws[f"{c}{r}"].value for c in "DEHJKLNPSTUY"}


def test_cliente_mixto_one_stop_adicion_genera_dos_op_y_un_contrato(entorno):
    eml = crear_correo(entorno, "mixto.eml", "CENTRO // DEALERS // ONE STOP // 500000001 // EMPRESA // 3 LINEAS // ADICIONES",
                       [EQUIPO, SIM1, SIM2], cuerpo="Pendiente\n\nReciben: Ana Lopez Ruiz y Luis Diaz Mora\n")
    cli.procesar(str(eml), None, None, FECHA, sufijo="ADICIONES")

    salida = config.SALIDA_DIR / "500000001_EMPRESA_PRUEBA_SA_DE_CV_ADICIONES"
    term = config.CONTRATOS_TERMINADOS_DIR / salida.name
    esperados = ["OP EQUIPOS.pdf", "OP EQUIPOS.xlsx", "OP SIM.pdf", "OP SIM.xlsx", "contrato_borrador.pdf"]
    assert sorted(p.name for p in term.iterdir()) == sorted(esperados)
    assert not (salida / "op_borrador.pdf").exists() and (salida / "revision.html").exists()

    # ---- contrato único (mismo para las dos OP), ADICIÓN, vendedor ONE STOP, personas del correo
    t = texto(salida / "contrato_borrador.pdf")
    assert "ALEJANDRO CASTERA ARELLANOS" in t and "ONE STOP MARKET AC425E" in t and "SERGIO YAHIR" not in t
    assert "ANA LOPEZ RUIZ Y LUIS DIAZ MORA" in t and "JUAN PEREZ LOPEZ" in t and "Sr." not in t
    assert "24 09 2026" in t and "08 10 2026" in t
    pg = fitz.open(salida / "contrato_borrador.pdf")[0]
    xs = [w for w in pg.get_text("words") if w[4] == "X" and 80 <= w[1] <= 96]
    assert len(xs) == 1 and abs(xs[0][0] - 131.9) < 1.5           # casilla ADICIÓN

    # ---- OP EQUIPOS: solo la línea con equipo, con precio/MPE
    eq = openpyxl.load_workbook(salida / "OP EQUIPOS.xlsx").active
    f = fila(eq, 16)
    assert (f["D"], f["J"], f["K"], f["N"]) == ("CDMX", 19999, 13199, 283.34) and int(f["L"]) == 24
    assert eq["D17"].value in (None, "") and eq["G12"].value == "240729MKT0305"
    assert eq["X6"].value.date() == FECHA
    assert eq["N12"].value == "INE //1234567890"

    # ---- OP SIM: 2 SIM, todo lo que no aplica en N/A, addon CTRL desde 'Addon Extra' del SAE
    sim = openpyxl.load_workbook(salida / "OP SIM.xlsx").active
    for r in (16, 17):
        f = fila(sim, r)
        assert (f["D"], f["J"], f["K"], f["N"]) == ("N/A", "N/A", "N/A", "N/A")
        assert (f["P"], f["S"], f["Y"]) == ("X", 50, 349)
    assert sim["D18"].value in (None, "") and sim["G12"].value == "240717MKT0205"   # SIM a 24 meses

    # ---- ni ShrinkToFit, ni formatos distintos entre renglones (archivos editables)
    for ws in (eq, sim):
        assert not [c.coordinate for row in ws.iter_rows(min_row=1, max_row=40) for c in row
                    if c.alignment and c.alignment.shrink_to_fit]
    assert len({(sim[f"C{r}"].font.sz, sim[f"C{r}"].alignment.horizontal) for r in (16, 17)}) == 1

    # ---- paquete de revisión: ambas OP, el INE y el aviso de canal
    h = (salida / "revision.html").read_text(encoding="utf-8")
    assert "OP EQUIPOS generada" in h and "OP SIM generada" in h and "ine_cotejo_p1.png" in h
    assert "Canal ONE STOP detectado" in h and "no se mezclan" in h
    assert all(fitz.open(p)[0].rect.width > 0 for p in salida.glob("*.pdf"))
    assert all(p.stat().st_size < config.LIMITE_PDF_KB * 1024 for p in salida.glob("*.pdf"))


def test_solo_contrato_no_toca_las_op_y_una_op_sola_borra_las_separadas(entorno):
    mixto = crear_correo(entorno, "mixto.eml", "CENTRO // DEALERS // FES PRIMECOMMS // 500000001 // RENOVACION",
                         [EQUIPO, SIM1])
    cli.procesar(str(mixto), None, None, FECHA)
    salida = config.SALIDA_DIR / "500000001_EMPRESA_PRUEBA_SA_DE_CV"
    term = config.CONTRATOS_TERMINADOS_DIR / salida.name
    assert (term / "OP EQUIPOS.xlsx").exists() and (term / "OP SIM.xlsx").exists()
    antes = {p.name: p.stat().st_mtime for p in salida.glob("OP *")}
    time.sleep(1.1)

    # 1) --solo-contrato: regenera el contrato y NO toca las OP (ni en salida/ ni en TERMINADOS)
    cli.procesar(str(mixto), None, None, FECHA, solo_contrato=True)
    assert {p.name: p.stat().st_mtime for p in salida.glob("OP *")} == antes
    assert (term / "OP EQUIPOS.xlsx").exists() and (term / "OP SIM.xlsx").exists()
    assert "SERGIO YAHIR VALDERRAMA VELOZ" in texto(term / "contrato_borrador.pdf")   # canal PRIME

    # 2) el mismo cliente ahora solo con equipos: queda UNA op_borrador y se borran las separadas
    solo_equipo = crear_correo(entorno, "equipo.eml", "CENTRO // DEALERS // FES PRIMECOMMS // 500000001 // RENOVACION",
                               [EQUIPO])
    cli.procesar(str(solo_equipo), None, None, FECHA)
    assert sorted(p.name for p in term.iterdir()) == ["contrato_borrador.pdf", "op_borrador.pdf", "op_borrador.xlsx"]
    assert not list(salida.glob("OP *")) and not list(salida.glob("op_equipos_p*.png"))


def test_equipo_desconocido_no_se_sustituye_y_queda_en_na_con_alerta(entorno):
    desconocido = ("5522222222", PLAN_EQ, "SAMSUNG GALAXY A14 128GB 4G", "NEGRO", 0)
    eml = crear_correo(entorno, "desc.eml", "CENTRO // DEALERS // FES PRIMECOMMS // 500000001 // RENOVACION", [desconocido])
    cli.procesar(str(eml), None, None, FECHA)
    salida = config.SALIDA_DIR / "500000001_EMPRESA_PRUEBA_SA_DE_CV"
    f = fila(openpyxl.load_workbook(salida / "op_borrador.xlsx").active, 16)
    assert f["J"] == "N/A" and f["N"] == "N/A"
    h = (salida / "revision.html").read_text(encoding="utf-8")
    assert "MISMO modelo" in h and "no se usa el precio de un equipo distinto" in h


def test_precio_de_lista_gana_cuando_el_master_difiere(entorno):
    zfold = ("5533333333", PLAN_EQ, "SAMSUNG GALAXY ZFOLD 8 ULTRA 512GB 5G", "NEGRO", 37759)
    eml = crear_correo(entorno, "z.eml", "CENTRO // DEALERS // FES PRIMECOMMS // 500000001 // RENOVACION", [zfold])
    cli.procesar(str(eml), None, None, FECHA)
    salida = config.SALIDA_DIR / "500000001_EMPRESA_PRUEBA_SA_DE_CV"
    f = fila(openpyxl.load_workbook(salida / "op_borrador.xlsx").active, 16)
    assert (f["J"], f["K"], f["N"]) == (53499, 37759, 655.84)          # lista $53,499, NO el Master $48,499
    assert "Consultar por qué difieren" in (salida / "revision.html").read_text(encoding="utf-8")
